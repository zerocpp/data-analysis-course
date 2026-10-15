"""Submission Validation v2 Conformance Tests（§39/§41：Same Rules 证明）。

目标（v2 §45：One Contract / One Single-Package Validator / Two Independent
Executions / One Server-only Batch Validator）：

1. Student Build PASS ⇔ Server Single-Package PASS（同一包，两端独立执行）；
2. 同一份错误语料在 Student 侧（submission-build / submission-validate）与
   Server 侧（submission-receive）必须得到**完全相同的 error_code** ——
   证明学生与助教对「一个包对不对」使用同一个程序；
3. Server 不信任学生本地 PASS：build 通过后手工删除成员 -> Server 重验即 FAIL；
4. Batch Validator：只有 Single-Package PASS Candidate 参加跨包检查，
   FAIL 包只标 NOT_RUN / REJECT；Excel 报告可读。

全部在 tmp_path 内完成：不依赖外部 corpus、不调用 LLM、不启动 subprocess。
"""
from __future__ import annotations

import json
import posixpath
import zipfile
from pathlib import Path

import pytest

from codesetarena.submission import contract as C
from codesetarena.submission.batch.excel_report import write_batch_xlsx
from codesetarena.submission.batch.validator import validate_batch
from codesetarena.submission.local import build_submission, validate_local_zip
from codesetarena.submission.receiver import receive_submission
from codesetarena.submission.safe_extract import safe_extract

TASK_STATEMENT = "data.csv 包含 region、amount 两列。请计算 amount 的总和并只输出整数。"
DIRECT_AGENT = 'def run(task, df, llm):\n    return int(df["amount"].sum())\n'
REFERENCE_SRC = 'def reference(df):\n    return int(df["amount"].sum())\n'
DATA_CSV = "region,amount\neast,10\nwest,20\neast,5\n"
ENV_EXAMPLE = "LLM_API_KEY=\nLLM_BASE_URL=\nLLM_MODEL=\n"
ENV_REAL = ("LLM_API_KEY=dummy-key\nLLM_BASE_URL=https://example.invalid/v1\n"
            "LLM_MODEL=dummy-model\nLLM_TIMEOUT=60\nLLM_TOKEN_BUDGET=10000\n")
STUDENT = "273001"
TASK_ID = "T_11111111111111111111111111111111"
PROTO = f"DataC1004_P1_{STUDENT}.zip"


# ---------------------------------------------------------------------------
# 构造工具（与 corpus 作者工具同构；仅用于本测试）
# ---------------------------------------------------------------------------


def make_workspace(root: Path, *, student_id: str = STUDENT, task_id: str = TASK_ID) -> Path:
    """扁平学生 workspace（run/export 网关输入，build_submission 的前置）。"""

    root.mkdir(parents=True, exist_ok=True)
    (root / ".env").write_text(ENV_REAL, encoding="utf-8")
    (root / ".env.example").write_text(ENV_EXAMPLE, encoding="utf-8")
    (root / "task.json").write_text(
        json.dumps({"task_id": task_id, "student_id": student_id,
                    "hypothesized_failure_mode": "FM02_SCHEMA",
                    "task_statement": TASK_STATEMENT}, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (root / "data.csv").write_text(DATA_CSV, encoding="utf-8")
    (root / "agent.py").write_text(DIRECT_AGENT, encoding="utf-8")
    (root / "reference.py").write_text(REFERENCE_SRC, encoding="utf-8")
    return root


def _apply_mutations(tree: Path, mutate: dict[str, object]) -> None:
    for rel, action in mutate.items():
        if action == "drop":
            target = tree / rel
            if target.is_dir():
                import shutil

                shutil.rmtree(target)
            elif target.exists():
                target.unlink()
        elif isinstance(action, dict):
            if action.get("write") is not None:
                target = tree / rel
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(action["write"], encoding="utf-8")
            elif action.get("json_merge"):
                target = tree / rel
                data = json.loads(target.read_text(encoding="utf-8"))
                data.update(action["json_merge"])
                target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                                  encoding="utf-8")
        else:
            raise ValueError(f"未知 mutate 动作: {rel!r}")


def make_submission_tree(
    root: Path,
    *,
    student_id: str = STUDENT,
    task_id: str = TASK_ID,
    phase: str = "phase1",
    mutate: dict[str, object] | None = None,
) -> Path:
    """构造一棵合法 Base Contract 提交树（可直接 zip 成提交包）。"""

    (root / "task/data").mkdir(parents=True)
    (root / "agent").mkdir(parents=True)
    (root / "manifest.json").write_text(
        json.dumps({"schema_version": C.SCHEMA_VERSION, "course": "DataC1004",
                    "submission_phase": phase, "student_id": student_id, "task_id": task_id,
                    "build_status": "passed", "files": {}}, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (root / "task/task.json").write_text(
        json.dumps({"task_id": task_id, "student_id": student_id,
                    "hypothesized_failure_mode": "FM02_SCHEMA",
                    "task_statement": TASK_STATEMENT}, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (root / "task/data/data.csv").write_text(DATA_CSV, encoding="utf-8")
    (root / "agent/agent.py").write_text(DIRECT_AGENT, encoding="utf-8")
    (root / "agent/agent_manifest.json").write_text(
        json.dumps({"schema_version": C.SCHEMA_VERSION, "student_id": student_id,
                    "task_id": task_id, "entrypoint": "agent.py"}, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (root / "task/reference.py").write_text(REFERENCE_SRC, encoding="utf-8")
    (root / ".env.example").write_text(ENV_EXAMPLE, encoding="utf-8")
    if mutate:
        _apply_mutations(root, mutate)
    return root


def write_zip(tree: Path, out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(p for p in tree.rglob("*") if p.is_file()):
            bundle.writestr(posixpath.join(*path.relative_to(tree).parts), path.read_bytes())


def _roster(tmp_path: Path) -> Path:
    roster = tmp_path / "roster.csv"
    roster.write_text(f"student_id\n{STUDENT}\n", encoding="utf-8")
    return roster


# ---------------------------------------------------------------------------
# §41：Student Build PASS ⇔ Server PASS；Server 不信任学生 PASS
# ---------------------------------------------------------------------------


def test_build_validate_receive_all_pass(tmp_path):
    """Build（本地校验 PASS）-> Local Valid -> Server Receive 全链路通过。"""

    record = build_submission(
        make_workspace(tmp_path / "ws"), phase="phase1", student_id=STUDENT,
        output_path=tmp_path / PROTO,
    )
    assert record["exit_code"] == 0
    assert record["local_validation"]["status"] == "PASS"
    pkg = Path(record["package_path"])
    assert pkg.name == PROTO and pkg.is_file()

    stu = validate_local_zip(pkg)
    assert stu["status"] == "PASS"

    result = receive_submission(pkg, tmp_path / "pool", _roster(tmp_path))
    assert result["status"] == "accepted", result
    assert result["error_code"] is None


def test_server_does_not_trust_student_pass(tmp_path):
    """学生 Build PASS 后手工篡改包：本地 Valid 与 Server Receive 都必须 FAIL。"""

    record = build_submission(
        make_workspace(tmp_path / "ws"), phase="phase1", student_id=STUDENT,
        output_path=tmp_path / "orig" / PROTO,
    )
    assert record["exit_code"] == 0 and record["local_validation"]["status"] == "PASS"

    extracted = tmp_path / "extract"
    safe_extract(Path(record["package_path"]), extracted)
    (extracted / "task/task.json").unlink()
    tampered = tmp_path / PROTO
    write_zip(extracted, tampered)

    stu = validate_local_zip(tampered)
    assert stu["status"] == "FAIL"
    assert stu["error_code"] == C.MISSING_TASK_JSON

    result = receive_submission(tampered, tmp_path / "pool", _roster(tmp_path))
    assert result["status"] == "rejected"
    assert result["error_code"] == C.MISSING_TASK_JSON
    assert not (tmp_path / "pool" / C.SUBMISSIONS / STUDENT).exists()  # 不写 canonical


# ---------------------------------------------------------------------------
# §39：错误语料 —— Student 与 Server 必须得到相同 error_code
# ---------------------------------------------------------------------------

CONFORMANCE_CASES = [
    ("missing_manifest", {"manifest.json": "drop"}, C.MISSING_MANIFEST),
    ("malformed_task_json", {"task/task.json": {"write": "{broken"}}, C.MALFORMED_JSON),
    ("invalid_fm",
     {"task/task.json": {"json_merge": {"hypothesized_failure_mode": "FM99_NOPE"}}},
     C.INVALID_FM),
    ("missing_data", {"task/data/data.csv": "drop"}, C.MISSING_TASK_DATA),
    ("missing_agent", {"agent/agent.py": "drop"}, C.MISSING_AGENT),
    ("missing_agent_manifest", {"agent/agent_manifest.json": "drop"}, C.MISSING_AGENT_MANIFEST),
    ("student_mismatch",
     {"manifest.json": {"json_merge": {"student_id": "273002"}},
      "task/task.json": {"json_merge": {"student_id": "273002"}},
      "agent/agent_manifest.json": {"json_merge": {"student_id": "273002"}}},
     C.STUDENT_ID_MISMATCH),
]


@pytest.mark.parametrize("name, mutate, expected", CONFORMANCE_CASES,
                         ids=[case[0] for case in CONFORMANCE_CASES])
def test_single_package_rules_identical_both_sides(tmp_path, name, mutate, expected):
    """每个错误语料：Student Valid 与 Server Receive 的 error_code 完全一致。"""

    pkg = tmp_path / PROTO
    write_zip(make_submission_tree(tmp_path / "tree", mutate=mutate), pkg)

    stu = validate_local_zip(pkg)
    assert stu["status"] == "FAIL"
    assert stu["error_code"] == expected, stu["errors"]

    result = receive_submission(pkg, tmp_path / "pool", _roster(tmp_path))
    assert result["status"] == "rejected"
    assert result["error_code"] == expected, result


def test_archive_name_rules_identical_both_sides(tmp_path):
    """非协议包名在两端都必须报 INVALID_ARCHIVE_NAME（同一词汇）。"""

    bad_name = tmp_path / "my_upload.zip"
    write_zip(make_submission_tree(tmp_path / "tree"), bad_name)

    stu = validate_local_zip(bad_name)
    assert stu["status"] == "FAIL"
    assert stu["error_code"] == C.INVALID_ARCHIVE_NAME

    result = receive_submission(bad_name, tmp_path / "pool", _roster(tmp_path))
    assert result["status"] == "rejected"
    assert result["error_code"] == C.INVALID_ARCHIVE_NAME


def test_phase_mismatch_rules_identical_both_sides(tmp_path):
    """内层 phase2 配外层 P1 包名：两端都必须报 SUBMISSION_PHASE_MISMATCH。"""

    pkg = tmp_path / PROTO
    write_zip(make_submission_tree(tmp_path / "tree", phase="phase2"), pkg)

    stu = validate_local_zip(pkg)
    assert stu["status"] == "FAIL"
    assert stu["error_code"] == C.SUBMISSION_PHASE_MISMATCH

    result = receive_submission(pkg, tmp_path / "pool", _roster(tmp_path))
    assert result["status"] == "rejected"
    assert result["error_code"] == C.SUBMISSION_PHASE_MISMATCH


# ---------------------------------------------------------------------------
# §24–§33：Batch Validator（第一遍单包 + 第二遍跨包 + Excel）
# ---------------------------------------------------------------------------


def test_batch_validator_single_then_cross(tmp_path):
    """Batch 第一遍只对 PASS Candidate 做跨包检查；FAIL 包 NOT_RUN / REJECT。"""

    batch_in = tmp_path / "batch_in"
    # A / B：单包 PASS 但共享 task_id -> 跨包 REJECT
    write_zip(make_submission_tree(tmp_path / "a", student_id="273001", task_id="T_1"),
              batch_in / "DataC1004_P1_273001.zip")
    write_zip(make_submission_tree(tmp_path / "b", student_id="273002", task_id="T_1"),
              batch_in / "DataC1004_P1_273002.zip")
    # C：单包 FAIL（缺 agent.py）
    bad = make_submission_tree(tmp_path / "c", student_id="273003", task_id="T_2")
    (bad / "agent/agent.py").unlink()
    write_zip(bad, batch_in / "DataC1004_P1_273003.zip")

    report = validate_batch(batch_in)
    assert report["status"] == "FAIL"
    assert report["summary"]["package_count"] == 3
    assert report["candidate_count"] == 2

    codes = [e["code"] for e in report["errors"]]
    assert C.BATCH_TASK_ID_DUPLICATE in codes
    dup = next(e for e in report["errors"] if e["code"] == C.BATCH_TASK_ID_DUPLICATE)
    assert set(dup["packages"]) == {"DataC1004_P1_273001.zip", "DataC1004_P1_273002.zip"}

    by_arch = {p["archive"]: p for p in report["packages"]}
    assert by_arch["DataC1004_P1_273001.zip"]["final_status"] == "REJECT"
    assert by_arch["DataC1004_P1_273002.zip"]["final_status"] == "REJECT"
    assert by_arch["DataC1004_P1_273003.zip"]["single_package_status"] == "FAIL"
    assert by_arch["DataC1004_P1_273003.zip"]["batch_status"] == "NOT_RUN"
    assert by_arch["DataC1004_P1_273003.zip"]["final_status"] == "REJECT"

    # 诊断信息包含所有可解析包中观察到的 task_id 重复（§27，不构成 reject）
    assert any(d["label"] == "task_id" for d in report["diagnostics"])


def test_batch_validator_clean_batch(tmp_path):
    """无冲突批次：全部 ACCEPTABLE，无跨包错误。"""

    batch_in = tmp_path / "batch_in"
    write_zip(make_submission_tree(tmp_path / "a", student_id="273001", task_id="T_1"),
              batch_in / "DataC1004_P1_273001.zip")
    write_zip(make_submission_tree(tmp_path / "b", student_id="273002", task_id="T_2"),
              batch_in / "DataC1004_P1_273002.zip")

    report = validate_batch(batch_in)
    assert report["status"] == "PASS"
    assert report["errors"] == []
    assert all(p["final_status"] == "ACCEPTABLE" for p in report["packages"])


def test_batch_validator_excel_report(tmp_path):
    """Excel 报告三个 Sheet 可读，主表行反映真实状态。"""

    batch_in = tmp_path / "batch_in"
    write_zip(make_submission_tree(tmp_path / "a", student_id="273001", task_id="T_1"),
              batch_in / "DataC1004_P1_273001.zip")
    report = validate_batch(batch_in)

    xlsx = write_batch_xlsx(report, tmp_path / "out")
    assert xlsx.name == "validation_report.xlsx" and xlsx.is_file()

    from openpyxl import load_workbook

    workbook = load_workbook(xlsx)
    assert workbook.sheetnames == ["Summary", "PackageResults", "BatchConflicts"]
    rows = list(workbook["PackageResults"].iter_rows(values_only=True))
    assert rows[0][0] == "Package"
    assert rows[1][0] == "DataC1004_P1_273001.zip"
    assert rows[1][4] == "PASS"   # Single Package
    assert rows[1][5] == "PASS"   # Batch
    assert rows[1][6] == "PASS"   # Final