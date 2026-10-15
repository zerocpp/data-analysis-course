"""Assignment 管理端 Excel 报告（§54–§57）。

机器可读结果一律 JSON（run.json / verifier.json / audit.json / leakage_scan.json）；
Excel 只是管理端人工视图，后续流程禁止解析 Excel。
"""

from __future__ import annotations

import json
from pathlib import Path


def _load(repo: Path, run_dir: Path, name: str) -> dict:
    path = run_dir / name
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def write_assignment_excel(repo_root, run_id: str) -> Path:
    from openpyxl import Workbook

    repo = Path(repo_root)
    run_dir = repo / "assignments" / run_id
    run = _load(repo, run_dir, "run.json")
    audit = _load(repo, run_dir, "audit.json")
    dist_check = _load(repo, run_dir, "distribution_verification.json")
    leakage = _load(repo, run_dir, "leakage_scan.json")

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    summary_fields = [
        ("assignment_run_id", run.get("assignment_run_id")),
        ("phase", run.get("phase")),
        ("method", run.get("method")),
        ("method_version", run.get("method_version")),
        ("policy_version", run.get("policy_version")),
        ("seed", run.get("seed")),
        ("r", run.get("r")),
        ("N (reviewers)", run.get("num_reviewers")),
        ("total_assignments", run.get("total_assignments")),
        ("verifier_status", run.get("verifier_status")),
        ("min_reviewer_load", run.get("min_reviewer_load")),
        ("max_reviewer_load", run.get("max_reviewer_load")),
        ("min_task_coverage", run.get("min_task_coverage")),
        ("max_task_coverage", run.get("max_task_coverage")),
        ("self_reviews", run.get("self_review_count")),
        ("duplicate_pairs", run.get("duplicate_pair_count")),
        ("diversity_score", run.get("diversity_score")),
        ("exposure_mae", run.get("exposure_mae")),
        ("distribution_status", run.get("distribution_status")),
        ("status", run.get("status")),
        ("input_snapshot_hash", run.get("input_snapshot_hash")),
    ]
    for key, value in summary_fields:
        ws.append([key, value])

    ws = wb.create_sheet("ReviewerLoads")
    ws.append(["reviewer_package_id", "student_id", "task_count", "reviewer_diversity",
               "exposure_col", "category_count"])
    snapshot = _load(repo, run_dir, "input_snapshot.json")
    reviewer_order = snapshot.get("reviewer_order", [])
    d_i = audit.get("d_i", [])
    n = len(reviewer_order)
    import numpy as np

    matrix = np.load(run_dir / "assignment.npy").tolist() if (run_dir / "assignment.npy").is_file() else None
    for i, reviewer in enumerate(reviewer_order):
        load = sum(matrix[i]) if matrix else 0
        cats = {snapshot["category_vector"][j] for j in range(n) if matrix and matrix[i][j]}
        ws.append([
            f"R{i + 1:06d}", reviewer, load,
            round(d_i[i], 4) if i < len(d_i) else "",
            None, len(cats),
        ])

    ws = wb.create_sheet("TaskCoverage")
    ws.append(["task_id", "owner_student_id", "category_index", "category_code", "reviewer_count"])
    if matrix is not None:
        task_order = snapshot["task_order"]
        owners = {t["task_id"]: t["owner_student_id"] for t in snapshot.get("tasks", [])}
        codes = {t["task_id"]: t["category_code"] for t in snapshot.get("tasks", [])}
        idxs = {t["task_id"]: t["category_index"] for t in snapshot.get("tasks", [])}
        col_sums = [sum(matrix[i][j] for i in range(n)) for j in range(n)]
        for j, task_id in enumerate(task_order):
            ws.append([task_id, owners.get(task_id, ""), idxs.get(task_id, ""),
                       codes.get(task_id, ""), col_sums[j]])

    ws = wb.create_sheet("CategoryDiversity")
    ws.append(["metric", "value"])
    for key in ("diversity_score", "d_case", "collisions_total", "exposure_mae"):
        ws.append([key, audit.get(key)])
    if "reviewer_diversity" in audit:
        for key, value in audit["reviewer_diversity"].items():
            ws.append([f"reviewer_diversity.{key}", value])

    ws = wb.create_sheet("PrivateAssignments")
    ws.append(["assignment_run_id", "phase", "reviewer_student_id", "reviewer_package_id",
               "task_id", "task_owner_student_id", "anonymous_task_id", "review_index"])
    registry_path = repo / "registry" / "assignments.jsonl"
    if registry_path.is_file():
        for line in registry_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            if row.get("assignment_run_id") == run_id:
                ws.append([
                    row.get("assignment_run_id"), row.get("phase"),
                    row.get("reviewer_student_id"), row.get("reviewer_package_id"),
                    row.get("task_id"), row.get("task_owner_student_id"),
                    row.get("anonymous_task_id"), row.get("review_index"),
                ])

    ws = wb.create_sheet("DistributionStatus")
    ws.append(["check", "value"])
    ws.append(["distribution_status", run.get("distribution_status")])
    ws.append(["packages_built", (dist_check.get("checks") or {}).get("packages_built")])
    ws.append(["identity_leaks", 0 if not leakage.get("identity_leaks") else len(leakage["identity_leaks"])])
    ws.append(["private_content_leaks",
               0 if not leakage.get("private_content_leaks") else len(leakage["private_content_leaks"])])
    if dist_check.get("errors"):
        ws.append(["distribution_errors", "; ".join(dist_check["errors"][:10])])

    out = run_dir / "assignment_report.xlsx"
    wb.save(out)
    return out